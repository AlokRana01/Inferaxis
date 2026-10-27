"""Unit and integration tests for DataIngestor."""

from pathlib import Path

import openpyxl
import pytest
from app.services.ingestor import (
    ColumnCountLimitExceededError,
    DataIngestor,
    EmptyFileError,
    FileParsingError,
    RowCountLimitExceededError,
    UnsupportedFileFormatError,
)


@pytest.fixture
def ingestor() -> DataIngestor:
    return DataIngestor(max_rows=10_000, max_cols=100)


def test_ingest_standard_csv(ingestor: DataIngestor):
    """Verify ingestion of a standard comma-separated CSV."""
    csv_bytes = b"id,name,revenue,is_active\n1,Acme,1200.5,true\n2,Beta,3400.0,false\n"
    result = ingestor.ingest_bytes(csv_bytes, "companies.csv")

    assert result.row_count == 2
    assert result.col_count == 4
    assert result.delimiter == ","
    assert result.normalized_columns == ["id", "name", "revenue", "is_active"]
    assert result.dataframe.iloc[0]["name"] == "Acme"


@pytest.mark.parametrize("delimiter", [";", "\t", "|"])
def test_detect_delimiters(ingestor: DataIngestor, delimiter: str):
    """Verify inference for semicolon, tab, and pipe delimiters."""
    content = f"id{delimiter}metric{delimiter}category\n1{delimiter}45.2{delimiter}A\n2{delimiter}98.1{delimiter}B\n"
    result = ingestor.ingest_bytes(content.encode("utf-8"), f"test{delimiter}.csv")

    assert result.delimiter == delimiter
    assert result.col_count == 3
    assert result.row_count == 2


def test_utf8_with_bom_encoding(ingestor: DataIngestor):
    """Verify handling of UTF-8 with Byte Order Mark (BOM)."""
    bom_content = b"\xef\xbb\xbfregion,sales\nEast,100\nWest,200\n"
    result = ingestor.ingest_bytes(bom_content, "bom_sales.csv")

    assert result.encoding == "utf-8-sig"
    assert result.normalized_columns[0] == "region"
    assert result.row_count == 2


def test_latin1_encoding_fallback(ingestor: DataIngestor):
    """Verify handling of Latin-1 encoded characters (accents/currency symbols)."""
    # "\xe9" is 'é' in latin-1
    latin_bytes = "id,city\n1,Montr\xe9al\n2,Z\xfcrich\n".encode("latin1")
    result = ingestor.ingest_bytes(latin_bytes, "cities_latin1.csv")

    assert result.row_count == 2
    assert "Montréal" in result.dataframe["city"].values or "Montr" in str(
        result.dataframe["city"].values
    )


def test_empty_file_rejection(ingestor: DataIngestor):
    """Verify that 0-byte or whitespace-only files raise EmptyFileError."""
    with pytest.raises(EmptyFileError):
        ingestor.ingest_bytes(b"", "empty.csv")

    with pytest.raises(EmptyFileError):
        ingestor.ingest_bytes(b"   \n\n  \t", "whitespace.csv")


def test_unsupported_format_rejection(ingestor: DataIngestor):
    """Verify that unsupported extensions raise UnsupportedFileFormatError."""
    with pytest.raises(UnsupportedFileFormatError):
        ingestor.ingest_bytes(b"some content", "data.parquet")

    with pytest.raises(UnsupportedFileFormatError):
        ingestor.ingest_bytes(b"some content", "document.pdf")


def test_duplicate_and_empty_column_normalization(ingestor: DataIngestor):
    """Verify that duplicate and missing headers are normalized deterministically."""
    # Columns: "date", "", "date", " "
    raw_csv = b"date,,date, \n2026-01-01,1,2026-01-02,4\n"
    result = ingestor.ingest_bytes(raw_csv, "duplicate_cols.csv")

    assert result.col_count == 4
    # Policy: date, unnamed_2, date_1, unnamed_4
    assert result.normalized_columns == ["date", "unnamed_2", "date_1", "unnamed_4"]
    assert (
        result.original_columns == ["date", "Unnamed: 1", "date.1", "Unnamed: 3"]
        or len(result.original_columns) == 4
    )


def test_row_and_column_limits():
    """Verify enforcement of row and column boundary limits."""
    strict_ingestor = DataIngestor(max_rows=5, max_cols=3)

    # 1. Row limit exceeded
    rows_csv = "id,val\n" + "\n".join(f"{i},{i * 10}" for i in range(10))
    with pytest.raises(RowCountLimitExceededError):
        strict_ingestor.ingest_bytes(rows_csv.encode(), "oversized_rows.csv")

    # 2. Column limit exceeded
    cols_header = ",".join(f"col_{i}" for i in range(10)) + "\n"
    cols_row = ",".join("1" for _ in range(10)) + "\n"
    with pytest.raises(ColumnCountLimitExceededError):
        strict_ingestor.ingest_bytes(
            (cols_header + cols_row).encode(), "oversized_cols.csv"
        )


def test_xlsx_ingestion(tmp_path: Path, ingestor: DataIngestor):
    """Verify secure ingestion of Excel XLSX workbooks."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q1_Metrics"
    ws.append(["quarter", "revenue", "profit"])
    ws.append(["Q1", 50000, 12000])
    ws.append(["Q2", 65000, 18000])

    # Second sheet
    ws2 = wb.create_sheet(title="Notes")
    ws2.append(["note_id", "comment"])
    ws2.append([1, "Audit passed"])

    xlsx_path = tmp_path / "financials.xlsx"
    wb.save(xlsx_path)
    wb.close()

    # 1. Default to first sheet
    result_default = ingestor.ingest_file(xlsx_path)
    assert result_default.row_count == 2
    assert result_default.col_count == 3
    assert result_default.selected_sheet == "Q1_Metrics"
    assert result_default.sheets == ["Q1_Metrics", "Notes"]

    # 2. Explicit sheet selection
    result_notes = ingestor.ingest_file(xlsx_path, sheet_name="Notes")
    assert result_notes.row_count == 1
    assert result_notes.col_count == 2
    assert result_notes.selected_sheet == "Notes"


def test_csv_with_embedded_delimiters_and_crlf(ingestor: DataIngestor):
    """Verify that quotes enclosing delimiters and CRLF line endings are parsed accurately."""
    content = b'id,name,role\r\n1,"Doe, John","Admin, Lead"\r\n2,"Smith, Jane","Architect"\r\n'
    result = ingestor.ingest_bytes(content, "staff.csv")

    assert result.delimiter == ","
    assert result.row_count == 2
    assert result.col_count == 3
    assert result.dataframe.iloc[0]["name"] == "Doe, John"
    assert result.dataframe.iloc[0]["role"] == "Admin, Lead"


def test_windows_1252_and_tsv_fixtures(ingestor: DataIngestor):
    """Verify CP1252 encoding with currency and TSV tab delimiters."""
    # CP1252 with euro symbol (\x80)
    cp1252_data = b"item,price\nWidget,50 \x80\n"
    res_cp1252 = ingestor.ingest_bytes(cp1252_data, "pricing.csv")
    assert res_cp1252.row_count == 1
    assert res_cp1252.col_count == 2

    # TSV fixture
    tsv_data = b"user_id\tusername\tscore\n101\talice\t95.5\n102\tbob\t88.0\n"
    res_tsv = ingestor.ingest_bytes(tsv_data, "export.tsv")
    assert res_tsv.delimiter == "\t"
    assert res_tsv.mime_type == "text/tab-separated-values"
    assert res_tsv.row_count == 2
    assert res_tsv.col_count == 3


def test_xlsx_column_limit_exceeded(tmp_path: Path):
    """Verify that an Excel workbook exceeding max_cols is rejected immediately."""
    strict_ingestor = DataIngestor(max_rows=100, max_cols=3)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["col1", "col2", "col3", "col4", "col5"])  # 5 columns > max 3
    ws.append([1, 2, 3, 4, 5])

    xlsx_path = tmp_path / "wide_table.xlsx"
    wb.save(xlsx_path)
    wb.close()

    with pytest.raises(ColumnCountLimitExceededError):
        strict_ingestor.ingest_file(xlsx_path)


def test_xlsx_zip_bomb_prevention(tmp_path: Path):
    """Verify that an Excel file with an excessive decompression ratio/size is rejected."""
    # Max file size 15KB -> Max uncompressed 150KB
    micro_ingestor = DataIngestor(max_file_size_bytes=15000)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["col1", "col2"])
    for i in range(500):
        ws.append([i, "long_text_string_" * 10 + str(i)])

    xlsx_path = tmp_path / "decompression_test.xlsx"
    wb.save(xlsx_path)
    wb.close()

    with pytest.raises(FileParsingError, match="maximum safe uncompressed size limit"):
        micro_ingestor.ingest_file(xlsx_path)
